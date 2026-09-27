from __future__ import annotations

import csv
import io
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

TEXT_EXTENSIONS = {".txt", ".md"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".csv", ".pdf", ".docx", ".xlsx", ".pptx", ".html", ".htm"}


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)


def _sections_from_text(text: str, max_sections: int = 80) -> list[dict[str, str]]:
    lines = [line.rstrip() for line in text.splitlines()]
    sections: list[dict[str, str]] = []
    current_heading = "Document"
    current: list[str] = []

    def flush() -> None:
        nonlocal current
        body = "\n".join(x for x in current if x.strip()).strip()
        if body:
            sections.append({"heading": current_heading, "text": body[:12000]})
        current = []

    for line in lines:
        stripped = line.strip()
        is_heading = (
            stripped.startswith("#")
            or (0 < len(stripped) <= 100 and stripped.isupper() and any(ch.isalpha() for ch in stripped))
        )
        if is_heading and len(sections) < max_sections:
            flush()
            current_heading = stripped.lstrip("# ") or "Section"
        else:
            current.append(line)
    flush()
    if not sections and text.strip():
        sections = [{"heading": "Document", "text": text[:12000]}]
    return sections[:max_sections]


def _payload(
    filename: str,
    suffix: str,
    text: str,
    engine: str,
    *,
    tables: list[dict[str, Any]] | None = None,
    pages: list[dict[str, Any]] | None = None,
    warnings: list[str] | None = None,
) -> dict[str, Any]:
    clean_text = text.strip()
    return {
        "fileName": filename,
        "fileType": suffix.lstrip(".") or "text",
        "processingEngine": engine,
        "status": "SUCCESS",
        "markdown": clean_text,
        "fullText": clean_text,
        "textPreview": clean_text[:4000],
        "characterCount": len(clean_text),
        "sections": _sections_from_text(clean_text),
        "tables": tables or [],
        "pages": pages or [],
        "warnings": warnings or [],
    }


def _parse_pdf(filename: str, content: bytes) -> dict[str, Any]:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("PyMuPDF is not installed. Run npm run dev:full again.") from exc

    pages: list[dict[str, Any]] = []
    text_parts: list[str] = []
    with fitz.open(stream=content, filetype="pdf") as document:
        for index, page in enumerate(document):
            text = page.get_text("text").strip()
            pages.append({"page": index + 1, "text": text[:30000]})
            if text:
                text_parts.append(f"## Page {index + 1}\n{text}")

    full_text = "\n\n".join(text_parts)
    warnings: list[str] = []
    if not full_text.strip():
        warnings.append("No selectable text was found. This appears to be a scanned/image PDF; OCR fallback is not enabled yet.")
    elif pages and len(full_text) / len(pages) < 60:
        warnings.append("Very little selectable text was extracted. Some pages may be scanned and could require OCR.")
    return _payload(filename, ".pdf", full_text, "PYMUPDF", pages=pages, warnings=warnings)


def _parse_docx(filename: str, content: bytes) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed. Run npm run dev:full again.") from exc

    document = Document(io.BytesIO(content))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    tables: list[dict[str, Any]] = []
    table_text: list[str] = []
    for index, table in enumerate(document.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        tables.append({"index": index + 1, "rows": rows[:500]})
        table_text.append(f"## Table {index + 1}\n" + "\n".join(" | ".join(row) for row in rows))
    text = "\n\n".join(paragraphs + table_text)
    return _payload(filename, ".docx", text, "PYTHON_DOCX", tables=tables)


def _parse_xlsx(filename: str, content: bytes) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is not installed. Run npm run dev:full again.") from exc

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    tables: list[dict[str, Any]] = []
    text_parts: list[str] = []
    for sheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
            if row_index >= 5000:
                break
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                rows.append(values)
        tables.append({"sheet": sheet.title, "rows": rows[:1000]})
        text_parts.append(f"## Sheet: {sheet.title}\n" + "\n".join(" | ".join(row) for row in rows))
    return _payload(filename, ".xlsx", "\n\n".join(text_parts), "OPENPYXL", tables=tables)


def _parse_pptx(filename: str, content: bytes) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx is not installed. Run npm run dev:full again.") from exc

    presentation = Presentation(io.BytesIO(content))
    pages: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    text_parts: list[str] = []
    for slide_index, slide in enumerate(presentation.slides):
        parts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "").strip()
            if text:
                parts.append(text)
            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                tables.append({"slide": slide_index + 1, "rows": rows})
                parts.extend(" | ".join(row) for row in rows)
        slide_text = "\n".join(parts)
        pages.append({"page": slide_index + 1, "text": slide_text[:30000]})
        if slide_text:
            text_parts.append(f"## Slide {slide_index + 1}\n{slide_text}")
    return _payload(filename, ".pptx", "\n\n".join(text_parts), "PYTHON_PPTX", tables=tables, pages=pages)


def _parse_html(filename: str, content: bytes, suffix: str) -> dict[str, Any]:
    decoded = content.decode("utf-8", errors="replace")
    extractor = _HTMLTextExtractor()
    extractor.feed(decoded)
    return _payload(filename, suffix, "\n".join(extractor.parts), "NATIVE_HTML")


def parse_document_bytes(filename: str, content: bytes) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")
    if suffix in TEXT_EXTENSIONS:
        return _payload(filename, suffix, content.decode("utf-8", errors="replace"), "NATIVE_TEXT")
    if suffix == ".csv":
        decoded = content.decode("utf-8", errors="replace")
        rows = list(csv.reader(io.StringIO(decoded)))
        text = "\n".join(" | ".join(row) for row in rows)
        return _payload(filename, suffix, text, "NATIVE_CSV", tables=[{"rows": rows[:1000]}])
    if suffix == ".pdf":
        return _parse_pdf(filename, content)
    if suffix == ".docx":
        return _parse_docx(filename, content)
    if suffix == ".xlsx":
        return _parse_xlsx(filename, content)
    if suffix == ".pptx":
        return _parse_pptx(filename, content)
    return _parse_html(filename, content, suffix)
